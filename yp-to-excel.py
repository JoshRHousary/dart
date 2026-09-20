"""Turn yp-taxonomy.json into an Excel workbook plus CSV exports.

Usage:
    python yp-to-excel.py [yp-taxonomy.json]
"""
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GOLD = PatternFill("solid", fgColor="D4AF37")
HEADER = Font(bold=True, color="1A1A1A")


def write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = GOLD
        cell.font = HEADER
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(list(row))
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows] or [0])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 4, 52)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    here = Path(__file__).parent
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else here / "yp-taxonomy.json"
    data = json.loads(src.read_text(encoding="utf-8"))

    cats = data["categories"]
    leaves = [c for c in cats if c["level"] == "category"]
    sections = [c for c in cats if c["level"] == "section"]
    subs = [c for c in cats if c["level"] == "subsection"]

    cat_rows = [(c["name"], c["section"], c["subsection"], c["id"])
                for c in sorted(leaves, key=lambda c: (c["section"], c["subsection"], c["name"]))]
    city_rows = sorted({(l["city"], l["province"]) for l in data["locations"]})

    per_section = Counter(c["section"] for c in leaves)
    sec_rows = [(s["name"], per_section.get(s["name"], 0), s["id"])
                for s in sorted(sections, key=lambda s: s["name"])]
    sub_rows = [(s["name"], s["section"], s["id"])
                for s in sorted(subs, key=lambda s: (s["section"], s["name"]))]
    prov_rows = sorted(Counter(l["province"] for l in data["locations"]).items())

    wb = Workbook()
    wb.remove(wb.active)

    info = wb.create_sheet("Source")
    for row in [
        ("Source", data["source"]),
        ("Scraped", data["scrapedAt"]),
        ("Cities", len(city_rows)),
        ("Categories", len(cat_rows)),
        ("Sections", len(sec_rows)),
        ("Subsections", len(sub_rows)),
    ]:
        info.append(row)
    info["A1"].font = HEADER
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 46

    write_sheet(wb, "Categories", ["Category (What)", "Section", "Subsection", "YP ID"], cat_rows)
    write_sheet(wb, "Cities", ["City (Where)", "Province"], city_rows)
    write_sheet(wb, "Sections", ["Section", "Categories", "YP ID"], sec_rows)
    write_sheet(wb, "Subsections", ["Subsection", "Section", "YP ID"], sub_rows)
    write_sheet(wb, "Provinces", ["Province", "Cities"], prov_rows)

    out = here / "yp-categories-cities.xlsx"
    wb.save(out)

    csv_dir = here / "yp-data-csv"
    csv_dir.mkdir(exist_ok=True)
    for ws in wb:
        with open(csv_dir / f"{ws.title}.csv", "w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerows(ws.iter_rows(values_only=True))

    print(f"Wrote {out}")
    print(f"  Categories {len(cat_rows)} | Cities {len(city_rows)} | "
          f"Sections {len(sec_rows)} | Subsections {len(sub_rows)}")


if __name__ == "__main__":
    main()
